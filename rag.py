"""RAG pipeline: load -> chunk -> embed -> FAISS (persisted) -> retrieve -> grounded answer with sources."""
import json
import re
import threading
import time
import warnings
import zipfile
from pathlib import Path

import docx
from docx.opc.exceptions import PackageNotFoundError
import pymupdf
from langchain_community.vectorstores import FAISS
from langchain_community.vectorstores.utils import DistanceStrategy
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_ollama import ChatOllama, OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic_settings import BaseSettings, SettingsConfigDict


class Settings(BaseSettings):
    model_config = SettingsConfigDict(env_file=".env", extra="ignore")
    ollama_url: str = "http://localhost:11434"
    embed_model: str = "nomic-embed-text"
    llm_model: str = "qwen2.5-coder:3b"
    index_dir: Path = Path("index")
    data_dir: Path = Path("data")
    chunk_size: int = 1000
    chunk_overlap: int = 150
    max_upload_mb: int = 25


settings = Settings()
SUPPORTED = {".pdf", ".docx", ".txt", ".md"}
NO_ANSWER = "I don't know based on the provided documents."
# Unit-normalized vectors + inner product = cosine similarity; scores are in [-1, 1], higher is better.
VS_KW = {"normalize_L2": True, "distance_strategy": DistanceStrategy.MAX_INNER_PRODUCT}
# LangChain warns this combo is unsupported, but it does normalize on add and search; the warning is noise.
warnings.filterwarnings("ignore", message="Normalizing L2 is not applicable")

PROMPT = ChatPromptTemplate.from_messages([
    ("system", "You answer questions using ONLY the numbered context passages below. "
               "Cite every claim with its passage number like [1]. "
               f"If the context does not contain the answer, reply exactly: \"{NO_ANSWER}\" "
               "Never use outside knowledge."),
    ("human", "Context:\n{context}\n\nQuestion: {question}"),
])


def _clean(text):
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", text)).strip()


def load(path: Path) -> list[Document]:
    """One Document per PDF page (keeps page numbers); one per file for DOCX/TXT/MD."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED)}")
    try:
        if ext == ".pdf":
            # Open from bytes: a failed open-by-path leaks the file handle on Windows (file then can't be deleted).
            with pymupdf.open(stream=path.read_bytes(), filetype="pdf") as doc:
                return [Document(_clean(p.get_text()), metadata={"source": path.name, "page": i})
                        for i, p in enumerate(doc, 1)]
        text = ("\n".join(p.text for p in docx.Document(path).paragraphs) if ext == ".docx"
                else path.read_text("utf-8", errors="replace"))
    except (pymupdf.FileDataError, zipfile.BadZipFile, PackageNotFoundError) as e:
        raise ValueError(f"{path.name} is corrupt or not a valid {ext} file ({e})") from e
    return [Document(_clean(text), metadata={"source": path.name, "page": None})]


class RAG:
    def __init__(self, index_dir=None, embeddings=None, llm=None):
        self.dir = Path(index_dir or settings.index_dir)
        self.emb = embeddings or OllamaEmbeddings(model=settings.embed_model, base_url=settings.ollama_url)
        self.llm = llm or ChatOllama(model=settings.llm_model, base_url=settings.ollama_url, temperature=0)
        self.chain = PROMPT | self.llm | StrOutputParser()
        self.splitter = RecursiveCharacterTextSplitter(chunk_size=settings.chunk_size,
                                                       chunk_overlap=settings.chunk_overlap)
        self.lock = threading.Lock()
        self.vs, self.meta = None, {}  # meta: {filename: [chunk ids]} -> enables re-upload as upsert
        if (self.dir / "index.faiss").exists():
            # index.pkl is written by this service only; never point index_dir at untrusted files.
            self.vs = FAISS.load_local(str(self.dir), self.emb, allow_dangerous_deserialization=True, **VS_KW)
            self.meta = json.loads((self.dir / "metadata.json").read_text("utf-8"))

    def ingest(self, path: Path) -> int:
        chunks = [c for c in self.splitter.split_documents(load(path)) if c.page_content.strip()]
        if not chunks:
            raise ValueError(f"No extractable text in {path.name} (scanned PDF? needs OCR).")
        ids = [f"{path.name}#{i}" for i in range(len(chunks))]
        for c, cid in zip(chunks, ids):
            c.metadata["chunk_id"] = cid
        with self.lock:
            if self.vs is not None and path.name in self.meta:
                self.vs.delete(self.meta.pop(path.name))
            if self.vs is None:
                self.vs = FAISS.from_documents(chunks, self.emb, ids=ids, **VS_KW)
            else:
                self.vs.add_documents(chunks, ids=ids)
            self.meta[path.name] = ids
            self.dir.mkdir(parents=True, exist_ok=True)
            self.vs.save_local(str(self.dir))
            (self.dir / "metadata.json").write_text(json.dumps(self.meta, indent=1), "utf-8")
        return len(chunks)

    def retrieve(self, query: str, k: int = 4) -> list[tuple[Document, float]]:
        if self.vs is None or not self.meta:
            return []
        return self.vs.similarity_search_with_score(query, k=k)

    def answer(self, query: str, k: int = 4) -> dict:
        t0 = time.perf_counter()
        hits = self.retrieve(query, k)
        answer = NO_ANSWER
        if hits:
            context = "\n\n".join(f"[{i}] ({d.metadata['source']}, page {d.metadata['page']})\n{d.page_content}"
                                  for i, (d, _) in enumerate(hits, 1))
            answer = self.chain.invoke({"context": context, "question": query}).strip()
        return {
            "answer": answer,
            "sources": [{"id": i, "file": d.metadata["source"], "page": d.metadata["page"],
                         "chunk_id": d.metadata["chunk_id"], "score": round(float(s), 4),
                         "snippet": d.page_content[:300]} for i, (d, s) in enumerate(hits, 1)],
            "latency_ms": round((time.perf_counter() - t0) * 1000),
        }

    def stats(self) -> dict:
        return {"documents": len(self.meta), "chunks": sum(map(len, self.meta.values())), "files": sorted(self.meta)}
